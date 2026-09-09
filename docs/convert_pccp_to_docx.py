"""Convert PCCP特征字典.md to a well-formatted Word document.

All LaTeX math is converted into *native Word equations* (OMML) via
latex2mathml -> MathML -> MML2OMML.XSL (Microsoft Office). The resulting
equations are fully editable inside Word / MathType (no images).
"""

import os
import re

from lxml import etree
from latex2mathml.converter import convert as latex_to_mathml

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

MD_PATH = r'E:\codes\ZZ-BK\docs\PCCP特征字典.md'
OUT_PATH = r'E:\codes\ZZ-BK\docs\PCCP特征字典.docx'
MML2OMML_XSL = r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'

OMML_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'

_xmlparser = etree.XMLParser(resolve_entities=False)
_omml_transform = etree.XSLT(etree.parse(MML2OMML_XSL))


# ============================== LaTeX -> OMML ==============================

def _preprocess(expr):
    """Normalise LaTeX not understood by latex2mathml into usable form."""
    def _cjk(m):
        return m.group(1)
    # 1. CJK text inside \text{}/\mathrm{}/\operatorname{} becomes plain chars
    expr = re.sub(
        r'\\(?:text|mathrm|operatorname|textnormal|mbox)\{'
        r'([^{}]*[\u3007\u3400-\u9fff\uff00-\uffef][^{}]*)\}', _cjk, expr)
    # 2. \text{...} -> \mathrm{...}
    expr = re.sub(r'\\text\{([^{}]*)\}', r'\\mathrm{\1}', expr)
    # 3. \operatorname{...} -> \mathrm{...}
    expr = re.sub(r'\\operatorname\{([^{}]*)\}', r'\\mathrm{\1}', expr)
    # 4. absolute / norm bars
    expr = expr.replace(r'\left\lvert', r'\left|').replace(r'\right\rvert', r'\right|')
    expr = expr.replace(r'\left\vert', r'\left|').replace(r'\right\vert', r'\right|')
    expr = expr.replace(r'\lvert', '|').replace(r'\rvert', '|')
    expr = expr.replace(r'\lVert', r'\|').replace(r'\rVert', r'\|')
    # 5. \mathbf1 -> \mathbf{1}
    expr = re.sub(r'\\mathbf([0-9A-Za-z])', r'\\mathbf{\1}', expr)
    # 6. \frac1K -> \frac{1}{K}
    expr = re.sub(r'\\frac(\d)([A-Za-z])', r'\\frac{\1}{\2}', expr)
    # 7. drop spacing commands unsupported by latex2mathml
    expr = re.sub(r'\\!\s*', '', expr)
    expr = expr.replace(r'\ ', ' ')
    return expr


def latex_to_omml(latex_str):
    """Return an lxml ``<m:oMath>`` element for the given LaTeX expression."""
    mathml = latex_to_mathml(_preprocess(latex_str))
    mml_root = etree.fromstring(mathml.encode('utf-8'), parser=_xmlparser)
    xslt_result = _omml_transform(mml_root)
    xml_str = str(xslt_result)
    xml_str = re.sub(r'^<\?xml[^>]*\?>', '', xml_str).strip()
    return etree.fromstring(xml_str.encode('utf-8'), parser=_xmlparser)


def add_omml(paragraph, omml_el, display=False):
    """Append an OMML element to a paragraph. display=True -> oMathPara."""
    if display:
        para = etree.SubElement(paragraph._p, f'{{{OMML_NS}}}oMathPara')
        jc = etree.SubElement(para, f'{{{OMML_NS}}}oMathParaPr')
        jc.set(f'{{{OMML_NS}}}jc', 'left')
        para.append(omml_el)
    else:
        paragraph._p.append(omml_el)


# ============================== Document helpers ============================

def set_cell_shading(cell, color_hex):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}" w:val="clear"/>'))


def set_cell_valign(cell, val='center'):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:vAlign {nsdecls("w")} w:val="{val}"/>'))


def set_table_borders(table, color='8EAADB', size='6'):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{size}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="single" w:sz="{size}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="{size}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="single" w:sz="{size}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="{size}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="single" w:sz="{size}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>')
    tblPr.append(borders)


def set_table_layout_fixed(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    if not tblPr.findall(qn('w:tblLayout')):
        tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))


def set_run_font(run, size_pt, name='微软雅黑', bold=False, color=None,
                 east_asia=True, mono=False):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    if mono:
        run.font.name = 'Consolas'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
    else:
        run.font.name = name
        if east_asia:
            run._element.rPr.rFonts.set(qn('w:eastAsia'), name)


def add_rich_text(paragraph, text, base_size=10, default_bold=False,
                  code_color=None):
    """Add paragraph content handling **bold**, `code`, and $math$ in order."""
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\$[^$]+\$)', text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, base_size, bold=True)
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, base_size, mono=True,
                         color=code_color or RGBColor(0x8B, 0x00, 0x00))
        elif part.startswith('$') and part.endswith('$') and len(part) > 2:
            try:
                add_omml(paragraph, latex_to_omml(part[1:-1]))
            except Exception:
                run = paragraph.add_run(part)
                set_run_font(run, base_size, bold=default_bold)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, base_size, bold=default_bold)


def add_formula_cell(paragraph, formula, base_size=9):
    """Add a formula column cell: native display equation, or plain text."""
    formula = formula.strip('$').strip()
    if re.search(r'\\[a-zA-Z]', formula) or re.search(r'[{}]|_|\^', formula):
        try:
            add_omml(paragraph, latex_to_omml(formula), display=True)
            return
        except Exception:
            pass
    paragraph.add_run(formula)


# ============================== Converter ===================================

def convert():
    with open(MD_PATH, 'r', encoding='utf-8') as f:
        content = f.read()
    lines = content.split('\n')

    doc = Document()
    doc.core_properties.title = 'PCCP 断丝识别特征字典'
    doc.core_properties.author = 'ZZ-BK'

    # ---------- page & base style ----------
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.3)
    sec.left_margin = sec.right_margin = Cm(2.5)

    normal = doc.styles['Normal']
    normal.font.name = '微软雅黑'
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for i, (size, color_hex) in enumerate([(18, '1F3864'), (14, '2E75B6'),
                                           (12, '2E75B6')], start=1):
        hs = doc.styles[f'Heading {i}']
        hs.font.name = '微软雅黑'
        hs.font.size = Pt(size)
        hs.font.bold = True
        hs.font.color.rgb = RGBColor.from_string(color_hex)
        hs._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        hs.paragraph_format.space_before = Pt(20 if i == 1 else 12)
        hs.paragraph_format.space_after = Pt(8)

    AVAIL_W = Cm(16.0)  # 21cm - 2*2.5cm margins

    i, n = 0, len(lines)
    while i < n:
        raw = lines[i]
        line = raw.rstrip()

        if not line.strip():
            i += 1
            continue

        # ---------- headings ----------
        if line.startswith('# ') and not line.startswith('## '):
            p = doc.add_heading(line[2:].strip(), level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        if line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue
        if line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue

        # ---------- tables ----------
        if line.strip().startswith('|'):
            tbl_rows = []
            while i < n and lines[i].strip().startswith('|'):
                tbl_rows.append(lines[i].strip())
                i += 1
            if len(tbl_rows) < 3:
                continue
            header = [c.strip() for c in tbl_rows[0].split('|')[1:-1]]
            data_rows = [[c.strip() for c in r.split('|')[1:-1]]
                         for r in tbl_rows[2:]]
            nc = len(header)
            nr = len(data_rows) + 1

            table = doc.add_table(rows=nr, cols=nc)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            set_table_borders(table)
            set_table_layout_fixed(table)

            if nc == 6:
                widths = [0.05, 0.16, 0.14, 0.16, 0.29, 0.20]
            else:
                widths = [1.0 / nc] * nc
            col_w = [AVAIL_W * w for w in widths]

            # header
            set_repeat_table_header(table.rows[0])
            for j, htxt in enumerate(header):
                cell = table.rows[0].cells[j]
                cell.width = col_w[j]
                set_cell_shading(cell, '1F3864')
                set_cell_valign(cell)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(htxt)
                set_run_font(run, 9, bold=True, color=RGBColor(255, 255, 255))

            # data
            for r_i, row in enumerate(data_rows):
                bg = 'EAF1F8' if r_i % 2 == 0 else 'FFFFFF'
                tr = table.rows[r_i + 1]
                for j in range(nc):
                    txt = row[j] if j < len(row) else ''
                    cell = tr.cells[j]
                    cell.width = col_w[j]
                    set_cell_shading(cell, bg)
                    set_cell_valign(cell)
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 \
                        else WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)

                    if j == 4:
                        add_formula_cell(p, txt)
                    else:
                        add_rich_text(p, txt, base_size=8)

            doc.add_paragraph()
            continue

        # ---------- bullets ----------
        if line.strip().startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(4)
            add_rich_text(p, line.strip()[2:], base_size=10)
            i += 1
            continue

        # ---------- plain paragraph ----------
        p = doc.add_paragraph()
        add_rich_text(p, line, base_size=10)
        i += 1

    doc.save(OUT_PATH)
    size_kb = os.path.getsize(OUT_PATH) / 1024
    print(f'OK -> {OUT_PATH}  ({size_kb:.0f} KB)')


if __name__ == '__main__':
    convert()